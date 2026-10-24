"""Data-backed bid readiness, company eligibility, quote costing, and proposal draft.

A draft is not a submitted bid; it cannot be a completed final bid without
verifiable notice conditions, company evidence, documented cost and approval.
"""
from __future__ import annotations
from decimal import Decimal,InvalidOperation,ROUND_HALF_UP
from io import BytesIO
from datetime import datetime,timezone
import json
from db import session
from g2b_case_engine import SINGLE_CASE,cost_summary,readiness_check

def _d(v):
    try:
        x=Decimal(str(v))
        if not x.is_finite() or x<0:raise InvalidOperation()
        return x
    except (ValueError,InvalidOperation,TypeError):
        return None

def evaluate_eligibility(bid_no=SINGLE_CASE):
    if bid_no!=SINGLE_CASE:raise ValueError("only selected tender")
    with session() as db:
        notice=db.execute("""SELECT raw FROM procureflow.tenders
                       WHERE source='g2b' AND external_id=%s""",(bid_no,)).fetchone()
        company=db.execute("""SELECT value FROM procureflow.g2b_case_decisions
                  WHERE bid_no=%s AND kind='company' AND key='registration'""",
                  (bid_no,)).fetchone()
        ev=db.execute("""SELECT evidence_type,verified_source,structured
                FROM procureflow.g2b_case_evidence WHERE bid_no=%s
                AND evidence_type IN ('company_registration','industry_codes',
                  'region_codes','requirements','qualification_evaluation')
                """,(bid_no,)).fetchall()
    raw=notice["raw"] if notice else None
    if not isinstance(raw,dict) or raw.get("bidNtceNo")!=bid_no:
        return {"status":"unknown","bid_no":bid_no,
                "reason":"official tender detail required","passed":False}
    if not company or not isinstance(company["value"],dict):
        return {"status":"unknown","bid_no":bid_no,
                "reason":"verified company registration and qualification evidence required",
                "passed":False}
    factual={r["evidence_type"] for r in ev if r["verified_source"]}
    if not {"company_registration","requirements","qualification_evaluation"}<=factual:
        return {"status":"unknown","bid_no":bid_no,
                "reason":"missing independently verified requirement or company evidence",
                "missing":sorted({"company_registration","requirements",
                                  "qualification_evaluation"}-factual),
                "passed":False}
    qualifications=next((r["structured"] for r in ev
          if r["evidence_type"]=="qualification_evaluation" and r["verified_source"]),{})
    result=qualifications.get("result")
    if result not in ("eligible","ineligible"):
        return {"status":"unknown","bid_no":bid_no,
                "reason":"qualification evaluation must give source-backed eligible/ineligible",
                "passed":False}
    return {"status":result,"passed":result=="eligible",
            "requirements_evidence_present":True,"bid_no":bid_no}

def documented_price_scenarios(bid_no=SINGLE_CASE):
    basic=cost_summary(bid_no)
    elig=evaluate_eligibility(bid_no)
    with session() as c:
        values=c.execute("""SELECT evidence_type,structured,verified_source
               FROM procureflow.g2b_case_evidence WHERE bid_no=%s
               AND evidence_type IN ('verified_base_amount','price_rules',
                    'cost_breakdown','supplier_quotes')""",(bid_no,)).fetchall()
    verified={v["evidence_type"]:v["structured"] for v in values
              if v["verified_source"] and isinstance(v["structured"],dict)}
    required={"verified_base_amount","price_rules","cost_breakdown","supplier_quotes"}
    missing=sorted(required-set(verified))
    if missing or not elig["passed"]:
        return {"status":"blocked","missing":missing,
                "eligibility":elig["status"],"current_documented_cost":basic,
                "scenarios":[],"recommended_bid_price":None}
    budget=_d(verified["verified_base_amount"].get("amount_krw"))
    floor=_d(verified["price_rules"].get("minimum_bid_ratio"))
    max_bid=_d(verified["price_rules"].get("maximum_bid_ratio"))
    costs=_d(verified["cost_breakdown"].get("total_krw"))
    vat=_d(verified["price_rules"].get("vat_ratio"))
    if None in (budget,floor,max_bid,costs,vat) or budget==0 or floor>max_bid or max_bid>1 or vat>1:
        return {"status":"blocked","missing":["price_formula_is_invalid"],
                "scenarios":[],"recommended_bid_price":None}
    scenario=[]
    for rate in (floor,(floor+max_bid)/2,max_bid):
        price=(budget*rate).quantize(Decimal("1"),rounding=ROUND_HALF_UP)
        cost_gross=(costs*(1+vat)).quantize(Decimal("1"),rounding=ROUND_HALF_UP)
        net=price-cost_gross
        scenario.append({"rate":str(rate),"bid_amount_krw":int(price),
                        "documented_cost_including_vat_krw":int(cost_gross),
                        "margin_krw":int(net),"economically_viable":net>=0})
    return {"status":"draft_computed_not_approved","scenarios":scenario,
            "recommended_bid_price":None,
            "requires_human_approval":True,"actual_submission":False}

def proposal_contents(bid_no=SINGLE_CASE):
    readiness=readiness_check(bid_no)
    with session() as db:
        details=db.execute("""SELECT title,agency,notice_date,deadline,method,
               raw FROM procureflow.tenders WHERE source='g2b' AND external_id=%s""",
               (bid_no,)).fetchone()
        docs=db.execute("""SELECT filename,sha256,extraction_state,url
                  FROM procureflow.g2b_case_documents WHERE bid_no=%s
                  ORDER BY filename LIMIT 80""",(bid_no,)).fetchall()
        ev=db.execute("""SELECT evidence_type,reference,summary,source_url,
                  verified_source FROM procureflow.g2b_case_evidence
                  WHERE bid_no=%s ORDER BY evidence_type,reference LIMIT 200""",
                  (bid_no,)).fetchall()
    return {"bid_no":bid_no,"title":details["title"] if details else "",
            "agency":details["agency"] if details else "",
            "date":details["notice_date"] if details else "",
            "deadline":details["deadline"] if details else "",
            "method":details["method"] if details else "",
            "stage":"draft_only_not_a_g2b_submission",
            "readiness":readiness,"eligibility":evaluate_eligibility(bid_no),
            "price":documented_price_scenarios(bid_no),
            "document_index":docs,"evidence_references":ev,
            "warnings":["Never assert company qualifications without evidence",
                        "Final bid forms must match original tender attachments",
                        "Manual review and bid participation remain separate"]}

def draft_docx(bid_no=SINGLE_CASE):
    """Downloadable true draft with explicit missing-evidence markers."""
    from docx import Document
    from docx.shared import Pt
    data=proposal_contents(bid_no)
    doc=Document()
    sec=doc.sections[0]
    doc.add_heading("입찰 사전검토 및 제출서류 준비서",0)
    doc.add_paragraph("본 문서는 제출된 입찰서가 아니라 내부 검토용 초안입니다.")
    table=doc.add_table(rows=0,cols=2);table.style="Table Grid"
    for name,key in [("공고번호","bid_no"),("공고명","title"),
                     ("발주기관","agency"),("공고일","date"),
                     ("마감일","deadline"),("계약방법","method")]:
        row=table.add_row().cells;row[0].text=name;row[1].text=str(data.get(key) or "미확보")
    for name,obj in [("입찰 참가자격 판정",data["eligibility"]),
                     ("원가 및 가격안",data["price"])]:
        doc.add_heading(name,level=1)
        doc.add_paragraph(json.dumps(obj,ensure_ascii=False,default=str,indent=2)[:7000])
    doc.add_heading("수집한 첨부 원본",level=1)
    if not data["document_index"]:doc.add_paragraph("다운로드한 첨부원본 없음.")
    for row in data["document_index"]:
        doc.add_paragraph(str(row["filename"])+
                          " / sha256="+str(row["sha256"])+
                          " / 상태="+str(row["extraction_state"]),style="List Bullet")
    doc.add_heading("확인해야 하는 미완료 요건",level=1)
    for group in data["readiness"]["phases"]:
        if not group["ready"]:
            doc.add_paragraph(group["key"]+": "+", ".join(group["missing_evidence"]),
                              style="List Bullet")
    doc.add_paragraph("투찰금액 및 참가자격 검증 완료 전에는 입찰참가·제출을 진행하지 않습니다.")
    stream=BytesIO();doc.save(stream)
    return stream.getvalue(),data
